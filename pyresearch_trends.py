from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, Series
import numpy
import glob

searchTerms = {
    "Brain Metastases": "Stereotactic AND (radiosurgery OR radiotherapy) AND (brain OR cranial OR cranium) AND (metastasis OR metastases)",
    "Spinal Metastases": "Stereotactic AND (radiosurgery OR radiotherapy) AND (spine OR brainstem OR spinal cord) AND (metastasis OR metastases)",
    "Meningioma": "Stereotactic AND (radiosurgery OR radiotherapy) AND Meningioma",
    "Glioblastoma": "Stereotactic AND (radiosurgery OR radiotherapy) AND (Glioblastoma OR GBM OR high grade glioma OR astrocytoma oligodendroglioma)",
    "AVM": "Stereotactic AND (radiosurgery OR radiotherapy) AND (Arteriovenous malformation OR AVM)",
    "Acoustic Neuroma": "Stereotactic AND (radiosurgery OR radiotherapy) AND (Acoustic neuroma OR vestibular schwanoma)",
    "Trigeminal Neuralgia": "Stereotactic AND (radiosurgery OR radiotherapy) AND (Trigeminal Neuralgia OR Tic Douloureux)",
    "Psychiatry": "Stereotactic AND (radiosurgery OR radiotherapy) AND (depression OR anxiety disorder OR obsessive compulsive OR obsessive compulsion)"
}
data = {k: {} for k, v in searchTerms.items()}


def outputOverviewData(file):
	subsections = numpy.load(file)
	subsections = subsections.item()
	del subsections['all']
	for section in subsections.keys():
		del subsections[section]['records']
		del subsections[section]['pubTypesData']
		del subsections[section]['pubSubsetsData']
		del subsections[section]['query']
	title = file[12:-16]
	data[title] = subsections


directory = 'subsections'
files = glob.glob("%s/*-subsections.npy" % directory)
print(files)
for f in files:
	outputOverviewData(f)

wb = Workbook()
ws = wb.active
ws.title = 'blank'
SRS = [
    {'gamma': 'Gamma Knife'},
    {'cyber': 'CyberKnife'},
    {'linac': 'LINAC'},
    {'novalis': 'Novalis'},
    {'tomo': 'TomoTherapy'}
]
indications = [
    'Brain Metastases',
    'Spinal Metastases',
    'Meningioma',
    'Glioblastoma',
    'AVM',
    'Acoustic Neuroma',
    'Trigeminal Neuralgia',
    'Psychiatry'
]
queries = {
    'gamma': 'AND (gamma knife OR gammaknife)',
    'cyber': 'AND (Cyberknife OR cyber knife)',
    'linac': 'AND linac',
    'novalis': 'AND Novalis',
    'tomo': 'AND TomoTherapy'
}
for modality in SRS:
    k = modality.keys()[0]
    v = modality.values()[0]
    ws = wb.create_sheet()
    ws.title = v
    print("Creating worksheet for %s" % v)
    ws.append([
    	v,
    	"%s %s" % (searchTerms[indications[0]], queries[k]),
    	"%s %s" % (searchTerms[indications[1]], queries[k]),
    	"%s %s" % (searchTerms[indications[2]], queries[k]),
    	"%s %s" % (searchTerms[indications[3]], queries[k]),
    	"%s %s" % (searchTerms[indications[4]], queries[k]),
    	"%s %s" % (searchTerms[indications[5]], queries[k]),
    	"%s %s" % (searchTerms[indications[6]], queries[k]),
    	"%s %s" % (searchTerms[indications[7]], queries[k])
    ])
    ws.append(['Years'] + indications)
    rowNum = 2
    chartStart = rowNum + 1
    oldestYear = min([
    	min([2016] + data[indications[0]][k]['yearlyData'].keys()),
    	min([2016] + data[indications[1]][k]['yearlyData'].keys()),
    	min([2016] + data[indications[2]][k]['yearlyData'].keys()),
    	min([2016] + data[indications[3]][k]['yearlyData'].keys()),
    	min([2016] + data[indications[4]][k]['yearlyData'].keys()),
    	min([2016] + data[indications[5]][k]['yearlyData'].keys()),
    	min([2016] + data[indications[6]][k]['yearlyData'].keys()),
    	min([2016] + data[indications[7]][k]['yearlyData'].keys()),
    ])
    for year in range(oldestYear, 2017):
        ws.append([
            year,
            data[indications[0]][k]['yearlyData'].get(year, 0),
            data[indications[1]][k]['yearlyData'].get(year, 0),
            data[indications[2]][k]['yearlyData'].get(year, 0),
            data[indications[3]][k]['yearlyData'].get(year, 0),
            data[indications[4]][k]['yearlyData'].get(year, 0),
            data[indications[5]][k]['yearlyData'].get(year, 0),
            data[indications[6]][k]['yearlyData'].get(year, 0),
            data[indications[7]][k]['yearlyData'].get(year, 0)
        ])
        rowNum += 1

    c1 = LineChart()
    c1.width = 30
    c1.height = 15
    dates = Reference(ws, min_row=chartStart, min_col=1, max_col=1, max_row=rowNum)
    s1 = Reference(ws, min_row=chartStart - 1, min_col=2, max_col=2, max_row=rowNum)
    s2 = Reference(ws, min_row=chartStart - 1, min_col=3, max_col=3, max_row=rowNum)
    s3 = Reference(ws, min_row=chartStart - 1, min_col=4, max_col=4, max_row=rowNum)
    s4 = Reference(ws, min_row=chartStart - 1, min_col=5, max_col=5, max_row=rowNum)
    s5 = Reference(ws, min_row=chartStart - 1, min_col=6, max_col=6, max_row=rowNum)
    s6 = Reference(ws, min_row=chartStart - 1, min_col=7, max_col=7, max_row=rowNum)
    s7 = Reference(ws, min_row=chartStart - 1, min_col=8, max_col=8, max_row=rowNum)
    s8 = Reference(ws, min_row=chartStart - 1, min_col=9, max_col=9, max_row=rowNum)
    c1.series.append(Series(s1, title_from_data=True))
    c1.series.append(Series(s2, title_from_data=True))
    c1.series.append(Series(s3, title_from_data=True))
    c1.series.append(Series(s4, title_from_data=True))
    c1.series.append(Series(s5, title_from_data=True))
    c1.series.append(Series(s6, title_from_data=True))
    c1.series.append(Series(s7, title_from_data=True))
    c1.series.append(Series(s8, title_from_data=True))
    c1.set_categories(dates)
    c1.title = "Articles per year for %s" % v
    c1.y_axis.title = 'Number of Articles'
    c1.x_axis.title = 'Year'
    ws.add_chart(c1, "A%d" % (rowNum + 5))

del wb['blank']
wb.save('%s/trends.xlsx' % directory)
